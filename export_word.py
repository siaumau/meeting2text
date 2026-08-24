import re
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

TABLE_SEP = re.compile(r"^:?-{3,}:?$")


def _set_east_asian_font(doc: Document, font_name: str) -> None:
    for style_name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3"):
        try:
            style = doc.styles[style_name]
        except KeyError:
            continue
        style.font.name = font_name
        rpr = style.element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = rpr.makeelement(qn("w:rFonts"), {})
            rpr.append(rfonts)
        rfonts.set(qn("w:eastAsia"), font_name)


def _add_runs(paragraph, text: str) -> None:
    parts = text.split("**")
    for idx, part in enumerate(parts):
        if not part:
            continue
        run = paragraph.add_run(part)
        if idx % 2 == 1:
            run.bold = True


def md_to_docx(md_text: str) -> Document:
    doc = Document()
    _set_east_asian_font(doc, "微軟正黑體")
    normal = doc.styles["Normal"]
    normal.font.size = Pt(11)

    lines = md_text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()

        if not line.strip():
            i += 1
            continue

        if line.lstrip().startswith("|"):
            rows: list[list[str]] = []
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(TABLE_SEP.match(c) for c in cells):
                    rows.append(cells)
                i += 1
            if rows:
                cols = max(len(r) for r in rows)
                table = doc.add_table(rows=len(rows), cols=cols)
                table.style = "Table Grid"
                for r, row in enumerate(rows):
                    for c in range(cols):
                        cell = table.rows[r].cells[c]
                        para = cell.paragraphs[0]
                        text = row[c] if c < len(row) else ""
                        _add_runs(para, text)
                        if r == 0:
                            for run in para.runs:
                                run.bold = True
            continue

        if line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=3)
        elif line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=2)
        elif line.startswith("# "):
            doc.add_heading(line[2:].strip(), level=1)
        elif re.match(r"^\s*[-*]\s+", line):
            para = doc.add_paragraph(style="List Bullet")
            _add_runs(para, re.sub(r"^\s*[-*]\s+", "", line))
        elif re.match(r"^\s*\d+[.、)]\s+", line):
            para = doc.add_paragraph(style="List Number")
            _add_runs(para, re.sub(r"^\s*\d+[.、)]\s+", "", line))
        else:
            para = doc.add_paragraph()
            _add_runs(para, line.strip())
        i += 1

    return doc


def export_docx(md_path: Path, docx_path: Path | None = None) -> Path:
    md_text = md_path.read_text(encoding="utf-8")
    target = docx_path or md_path.with_suffix(".docx")
    md_to_docx(md_text).save(str(target))
    return target


if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit("用法：python export_word.py <會議紀錄.md> [輸出.docx]")
    out = export_docx(
        Path(sys.argv[1]), Path(sys.argv[2]) if len(sys.argv) > 2 else None
    )
    print(f"已匯出：{out}")
